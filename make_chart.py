"""Render the README chart from the recalculated workbook (run after build + recalc)."""
import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt
from matplotlib.ticker import FuncFormatter
from openpyxl import load_workbook

wb = load_workbook("retirement_stress_test_model.xlsx", data_only=True)
s, p = wb["Scenarios"], wb["Projection"]

BLUE, ORANGE = "#2a78d6", "#eb6834"
INK, MUTED, SURF = "#0b0b0b", "#52514e", "#fcfcfb"
money = FuncFormatter(lambda v, _: f"-${abs(v)/1e6:.1f}M" if v < 0 else f"${v/1e6:.1f}M")

fig, (ax1, ax2) = plt.subplots(2, 1, figsize=(10, 9), facecolor=SURF)
for ax in (ax1, ax2):
    ax.set_facecolor(SURF)
    for side in ("top", "right"):
        ax.spines[side].set_visible(False)
    ax.spines["left"].set_color("#cccccc"); ax.spines["bottom"].set_color("#cccccc")
    ax.tick_params(colors=MUTED)
    ax.grid(axis="y", color="#e6e6e3", linewidth=0.8); ax.set_axisbelow(True)
    ax.yaxis.set_major_formatter(money)

# Panel 1: balance by age, selected scenario
ages = [p.cell(r, 2).value for r in range(6, 63)]
bal = [p.cell(r, 8).value for r in range(6, 63)]
ax1.plot(ages, bal, color=BLUE, linewidth=2)
ret_age = wb["Assumptions"]["B5"].value
ax1.axvline(ret_age, color=MUTED, linewidth=1, linestyle=":")
ax1.text(ret_age + 0.6, max(bal) * 0.15, "Retirement at 65", color=MUTED, fontsize=9)
ax1.set_title(f"Balanced portfolio with 5% TIPS: balance by age (expected case)",
              loc="left", color=INK, fontsize=12, fontweight="bold")
ax1.set_xlabel("Age", color=MUTED)

# Panel 2: expected vs bad-luck ending balance, 9 scenarios
labels = [f"{s.cell(r,1).value[:4]}.\n{int(round(s.cell(r,2).value*100))}% TIPS" for r in range(5, 14)]
exp_ = [s.cell(r, 11).value for r in range(5, 14)]
bad = [s.cell(r, 14).value for r in range(5, 14)]
x = range(len(labels)); w = 0.38
ax2.bar([i - w/2 - 0.01 for i in x], exp_, w, color=BLUE, label="Expected returns")
ax2.bar([i + w/2 + 0.01 for i in x], bad, w, color=ORANGE, label="Bad-luck case (~10th percentile)")
ax2.axhline(0, color=MUTED, linewidth=1)
ax2.set_xticks(list(x)); ax2.set_xticklabels(labels, fontsize=8.5, color=MUTED)
ax2.set_title("Money left at age 92: expected vs bad-luck returns",
              loc="left", color=INK, fontsize=12, fontweight="bold")
ax2.set_ylabel("Ending balance (below $0 = money runs out)", color=MUTED, fontsize=9)
ax2.legend(frameon=False, loc="upper left", labelcolor=INK)

fig.text(0.01, 0.005, "Illustrative assumptions for education only, not investment advice.",
         color=MUTED, fontsize=8)
fig.tight_layout(rect=(0, 0.02, 1, 1))
fig.savefig("model_overview.png", dpi=150, facecolor=SURF)
print("saved model_overview.png")

# ---------------------------------------------------------------- Account comparison chart
a = wb["Account Comparison"]
names = [a.cell(r, 1).value for r in range(26, 30)]
vals = [a.cell(r, 6).value for r in range(26, 30)]
rates = [a.cell(r, 1).value for r in range(35, 40)]
adv = [a.cell(r, 5).value for r in range(35, 40)]
kfmt = FuncFormatter(lambda v, _: f"-${abs(v)/1e3:,.0f}K" if v < 0 else f"${v/1e3:,.0f}K")

fig, (b1, b2) = plt.subplots(2, 1, figsize=(10, 8.5), facecolor=SURF,
                             gridspec_kw={"height_ratios": [1.1, 1]})
for ax in (b1, b2):
    ax.set_facecolor(SURF)
    for side in ("top", "right"):
        ax.spines[side].set_visible(False)
    ax.spines["left"].set_color("#cccccc"); ax.spines["bottom"].set_color("#cccccc")
    ax.tick_params(colors=MUTED)

y = list(range(len(names)))[::-1]
b1.barh(y, vals, height=0.55, color=BLUE)
b1.set_yticks(y); b1.set_yticklabels(names, color=INK)
b1.xaxis.set_major_formatter(kfmt)
b1.grid(axis="x", color="#e6e6e3", linewidth=0.8); b1.set_axisbelow(True)
for yi, v in zip(y, vals):
    b1.text(v + max(vals) * 0.01, yi, f"${v/1e3:,.0f}K", va="center", color=INK, fontsize=10)
b1.set_xlim(0, max(vals) * 1.15)
b1.set_title("After-tax value at 65 from $7,500/yr of pre-tax income",
             loc="left", color=INK, fontsize=12, fontweight="bold")

labels2 = [f"{r*100:.0f}%" for r in rates]
cols = [ORANGE if v > 0.5 else (BLUE if v < -0.5 else MUTED) for v in adv]
b2.bar(labels2, adv, width=0.55, color=cols)
b2.axhline(0, color=MUTED, linewidth=1)
b2.yaxis.set_major_formatter(kfmt)
b2.grid(axis="y", color="#e6e6e3", linewidth=0.8); b2.set_axisbelow(True)
b2.set_xlabel("Tax rate in retirement (current rate: 22%)", color=MUTED)
b2.set_ylabel("Roth minus Traditional ($)", color=MUTED, fontsize=9)
b2.set_title("Roth vs Traditional IRA: it depends on your future tax rate",
             loc="left", color=INK, fontsize=12, fontweight="bold")
from matplotlib.patches import Patch
b2.legend(handles=[Patch(color=BLUE, label="Traditional better"), Patch(color=ORANGE, label="Roth better")],
          frameon=False, loc="upper left", labelcolor=INK)
b2.text(2, 3000, "Tie", ha="center", color=MUTED, fontsize=9)

fig.text(0.01, 0.005, "Illustrative assumptions for education only, not tax or investment advice.",
         color=MUTED, fontsize=8)
fig.tight_layout(rect=(0, 0.02, 1, 1))
fig.savefig("account_comparison.png", dpi=150, facecolor=SURF)
print("saved account_comparison.png")

# ---------------------------------------------------------------- Business cash-flow chart
bw, sm = wb["Business Cash Flow"], wb["Simulation"]
AQUA, YELLOW_C = "#1baf7a", "#eda100"
months = list(range(1, 25))
fig, (c1, c2) = plt.subplots(2, 1, figsize=(10, 9), facecolor=SURF)
for ax in (c1, c2):
    ax.set_facecolor(SURF)
    for side in ("top", "right"):
        ax.spines[side].set_visible(False)
    ax.spines["left"].set_color("#cccccc"); ax.spines["bottom"].set_color("#cccccc")
    ax.tick_params(colors=MUTED)
    ax.grid(axis="y", color="#e6e6e3", linewidth=0.8); ax.set_axisbelow(True)
    ax.yaxis.set_major_formatter(kfmt)
    ax.axhline(0, color=MUTED, linewidth=1)
    ax.set_xlim(1, 27.5); ax.set_xlabel("Month", color=MUTED)

for si, colr in enumerate([BLUE, ORANGE, AQUA, YELLOW_C]):
    col = 2 + si * 7 + 4
    vals = [bw.cell(r, col).value for r in range(21, 45)]
    name = bw.cell(13 + si, 1).value
    c1.plot(months, vals, color=colr, linewidth=2, label=name)
    c1.text(24.4, vals[-1], name, color=INK, fontsize=9, va="center")
c1.legend(frameon=False, loc="lower left", labelcolor=INK, fontsize=9)
c1.set_title("Example Co. cash balance: four what-if scenarios",
             loc="left", color=INK, fontsize=12, fontweight="bold")

p = {k: [sm.cell(r, col).value for r in range(14, 38)] for k, col in
     zip(("p10", "p25", "p50", "p75", "p90"), range(2, 7))}
c2.fill_between(months, p["p10"], p["p90"], color=BLUE, alpha=0.15, linewidth=0, label="10th–90th percentile")
c2.fill_between(months, p["p25"], p["p75"], color=BLUE, alpha=0.30, linewidth=0, label="25th–75th percentile")
c2.plot(months, p["p50"], color=BLUE, linewidth=2, label="Median")
surv = sm["B6"].value
c2.set_title(f"1,000 simulated futures: {surv:.0%} chance cash stays positive for 24 months",
             loc="left", color=INK, fontsize=12, fontweight="bold")
c2.legend(frameon=False, loc="upper left", labelcolor=INK, fontsize=9)

fig.text(0.01, 0.005, "Fictional business with made-up numbers, for demonstration only.",
         color=MUTED, fontsize=8)
fig.tight_layout(rect=(0, 0.02, 1, 1))
fig.savefig("business_cashflow.png", dpi=150, facecolor=SURF)
print("saved business_cashflow.png")

# ---------------------------------------------------------------- Practice ROI chart
pr = wb["Practice ROI"]
AQUA = "#1baf7a"
fees = [[pr.cell(28, c).value for c in cols] for cols in ((2, 3, 4), (5, 6, 7))]
ann = [[pr.cell(29, c).value for c in cols] for cols in ((2, 3, 4), (5, 6, 7))]
ins = [[pr.cell(30, c).value + pr.cell(31, c).value for c in cols] for cols in ((2, 3, 4), (5, 6, 7))]
cost = pr["B19"].value
tot3 = (pr["B38"].value, pr["C38"].value)

fig, ax = plt.subplots(figsize=(10, 6), facecolor=SURF)
ax.set_facecolor(SURF)
for side in ("top", "right"):
    ax.spines[side].set_visible(False)
ax.spines["left"].set_color("#cccccc"); ax.spines["bottom"].set_color("#cccccc")
ax.tick_params(colors=MUTED); ax.grid(axis="y", color="#e6e6e3", linewidth=0.8); ax.set_axisbelow(True)
ax.yaxis.set_major_formatter(kfmt)
w = 0.36
for si, (off, lab) in enumerate(((-w / 2 - 0.02, "Standard"), (w / 2 + 0.02, "Data-driven"))):
    xs = [i + off for i in range(3)]
    b1 = ax.bar(xs, fees[si], w, color=BLUE, label="Advisory fees" if si == 0 else None, edgecolor=SURF, linewidth=1.5)
    b2 = ax.bar(xs, ann[si], w, bottom=fees[si], color=ORANGE, label="Annuity commissions" if si == 0 else None,
                edgecolor=SURF, linewidth=1.5)
    bottoms = [f + a for f, a in zip(fees[si], ann[si])]
    ax.bar(xs, ins[si], w, bottom=bottoms, color=AQUA, label="Insurance commissions" if si == 0 else None,
           edgecolor=SURF, linewidth=1.5)
    for x, f, a, n in zip(xs, fees[si], ann[si], ins[si]):
        ax.text(x, f + a + n + 1500, f"${(f + a + n)/1e3:,.0f}K", ha="center", color=INK, fontsize=9)
        ax.text(x, -6500, lab, ha="center", color=MUTED, fontsize=8)
ax.axhline(cost, color=INK, linewidth=1, linestyle="--")
ax.text(-0.45, cost + 1500, "Planner cost \\$45K/yr", color=INK, fontsize=9, ha="left")
ax.set_xticks(range(3)); ax.set_xticklabels(["Year 1", "Year 2", "Year 3"], color=INK)
ax.tick_params(axis="x", pad=18)
ax.set_ylim(-9000, max(f + a + n for f, a, n in zip(fees[1], ann[1], ins[1])) * 1.15)
ax.set_title(f"Revenue from one new planner: \\${tot3[0]/1e3:,.0f}K standard vs \\${tot3[1]/1e3:,.0f}K data-driven over 3 years",
             loc="left", color=INK, fontsize=12, fontweight="bold")
ax.legend(frameon=False, loc="upper left", labelcolor=INK, fontsize=9)
fig.text(0.01, 0.005, "Illustrative assumptions; replace with the firm's figures.", color=MUTED, fontsize=8)
fig.tight_layout(rect=(0, 0.02, 1, 1))
fig.savefig("practice_roi.png", dpi=150, facecolor=SURF)
print("saved practice_roi.png")
