"""
Business Cash-Flow Simulation (fictional small business, made-up numbers)
Adds two sheets to the workbook built by build_model.py:

  Business Cash Flow - 24-month driver-based model (customers x price - costs) for 4 what-if
                       scenarios, all live Excel formulas, with runway and breakeven
  Simulation         - 1,000-run Monte Carlo of monthly growth and churn: cash percentiles by
                       month and the probability the business stays cash-positive

Run after build_model.py:  python business_cashflow.py
"""
import numpy as np
from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.chart import LineChart, Reference

FILE = "retirement_stress_test_model.xlsx"
SEED, RUNS, MONTHS = 42, 1000, 24
GROWTH_SD, CHURN_SD, MONTH_SD = 0.015, 0.005, 0.01  # business-level and month-to-month uncertainty

F = "Arial"
BLUE = Font(name=F, color="0000FF"); BLACK = Font(name=F, color="000000")
BOLD = Font(name=F, bold=True); TITLE = Font(name=F, bold=True, size=14)
NOTE = Font(name=F, size=9, italic=True)
HDR = Font(name=F, bold=True, color="FFFFFF"); HDR_FILL = PatternFill("solid", fgColor="1F3864")
YELLOW = PatternFill("solid", fgColor="FFFF00")
USD = '$#,##0;($#,##0);"-"'; PCT = '0.0%;(0.0%);"-"'

def header(ws, row, labels, col=1):
    for i, t in enumerate(labels):
        c = ws.cell(row=row, column=col + i, value=t)
        c.font = HDR; c.fill = HDR_FILL; c.alignment = Alignment(horizontal="center", wrap_text=True)

# Inputs (fictional business)
inputs = {
    "cash": 300000, "customers": 250, "price": 150, "growth": 0.05,
    "churn": 0.02, "cogs": 0.25, "fixed": 45000,
}
scenarios = [  # name, price change, growth, churn, fixed-cost change
    ("Base plan", 0.00, inputs["growth"], inputs["churn"], 0.00),
    ("Raise price 10%", 0.10, inputs["growth"], inputs["churn"] + 0.005, 0.00),
    ("Cut fixed costs 10%", 0.00, inputs["growth"], inputs["churn"], -0.10),
    ("Downturn", 0.00, 0.01, 0.04, 0.00),
]

wb = load_workbook(FILE)
ws = wb.create_sheet("Business Cash Flow")
ws["A1"] = "Business Cash-Flow Model: Example Co. (fictional)"; ws["A1"].font = TITLE
ws["A2"] = ("Made-up small business for demonstration. Driver-based: customers x price, less costs. "
            "Blue = input, black = formula.")
ws["A2"].font = NOTE

header(ws, 4, ["Business inputs", "Value"])
rows = [("Starting cash ($)", "cash", USD), ("Starting customers", "customers", "0"),
        ("Monthly price per customer ($)", "price", USD),
        ("New customers per month (% of base)", "growth", PCT), ("Monthly churn (% lost)", "churn", PCT),
        ("Cost of service (% of revenue)", "cogs", PCT), ("Fixed operating costs per month ($)", "fixed", USD)]
cell = {}
for i, (lab, key, fmt) in enumerate(rows):
    r = 5 + i
    ws.cell(row=r, column=1, value=lab).font = BLACK
    c = ws.cell(row=r, column=2, value=inputs[key]); c.font = BLUE; c.number_format = fmt
    cell[key] = f"$B${r}"
# cash B5, customers B6, price B7, growth B8, churn B9, cogs B10, fixed B11

header(ws, 4, ["Scenario", "Price change", "New customers / mo", "Churn / mo", "Fixed-cost change"], col=4)
for i, (name, dp, g, ch, dfix) in enumerate(scenarios):
    r = 5 + i
    ws.cell(row=r, column=4, value=name).font = BOLD
    vals = [dp, g, ch, dfix]
    if i in (0, 1, 2):  # tie growth/churn to base inputs so changing B8/B9 flows through
        vals[1] = "=$B$8"
        vals[2] = "=$B$9+0.005" if i == 1 else "=$B$9"
    for j, v in enumerate(vals):
        c = ws.cell(row=r, column=5 + j, value=v); c.number_format = PCT
        c.font = BLACK if isinstance(v, str) else BLUE
ws["D10"] = "Raising price assumes slightly higher churn (+0.5 pt) as some customers leave."
ws["D10"].font = NOTE

# Summary block
header(ws, 12, ["Scenario", "Month-24 revenue ($)", "Month-24 monthly profit ($)",
                "Ending cash ($)", "Cash runs out in month", "Breakeven month"], col=1)

# Monthly projection: 7 columns per scenario starting at column A for month
TOP = 20  # header row
ws.cell(row=TOP - 1, column=1, value="Monthly projection").font = BOLD
ws.cell(row=TOP, column=1, value="Month").font = HDR; ws.cell(row=TOP, column=1).fill = HDR_FILL
sub = ["Customers", "Revenue ($)", "Costs ($)", "Profit / loss ($)", "Cash ($)", "neg flag", "BE flag"]
first, last = TOP + 1, TOP + MONTHS
for m in range(1, MONTHS + 1):
    ws.cell(row=TOP + m, column=1, value=m).font = BLACK
for si, (name, *_rest) in enumerate(scenarios):
    p = 5 + si                       # scenario parameter row
    c0 = 2 + si * 7                  # first column of block
    L = [ws.cell(row=1, column=c0 + k).column_letter for k in range(7)]
    ws.cell(row=TOP - 1, column=c0, value=name).font = BOLD
    header(ws, TOP, sub, col=c0)
    for m in range(1, MONTHS + 1):
        r = TOP + m
        prev_c = f"{L[0]}{r-1}" if m > 1 else "$B$6"
        prev_cash = f"{L[4]}{r-1}" if m > 1 else "$B$5"
        f = [
            f"={prev_c}*(1+$F${p}-$G${p})",
            f"={L[0]}{r}*$B$7*(1+$E${p})",
            f"={L[1]}{r}*$B$10+$B$11*(1+$H${p})",
            f"={L[1]}{r}-{L[2]}{r}",
            f"={prev_cash}+{L[3]}{r}",
            f"=IF({L[4]}{r}<0,$A{r},999)",
            f"=IF({L[3]}{r}>=0,$A{r},999)",
        ]
        for k, formula in enumerate(f):
            c = ws.cell(row=r, column=c0 + k, value=formula); c.font = BLACK
            c.number_format = "#,##0" if k == 0 else ("0" if k >= 5 else USD)
    # summary row
    sr = 13 + si
    ws.cell(row=sr, column=1, value=name).font = BOLD
    ws.cell(row=sr, column=2, value=f"={L[1]}{last}").number_format = USD
    ws.cell(row=sr, column=3, value=f"={L[3]}{last}").number_format = USD
    ws.cell(row=sr, column=4, value=f"={L[4]}{last}").number_format = USD
    ws.cell(row=sr, column=5, value=f'=IF(MIN({L[5]}{first}:{L[5]}{last})=999,"Not within 24",MIN({L[5]}{first}:{L[5]}{last}))')
    ws.cell(row=sr, column=6, value=f'=IF(MIN({L[6]}{first}:{L[6]}{last})=999,"Not within 24",MIN({L[6]}{first}:{L[6]}{last}))')
    for col in range(2, 7):
        ws.cell(row=sr, column=col).font = BLACK
        ws.cell(row=sr, column=col).alignment = Alignment(horizontal="right")
ws["A18"] = "Flag columns (neg / BE) mark the first month cash goes negative or profit turns positive (999 = never)."
ws["A18"].font = NOTE
ws.column_dimensions["A"].width = 38
for k in range(2, 2 + 7 * 4):
    ws.column_dimensions[ws.cell(row=1, column=k).column_letter].width = 14
ws.column_dimensions["D"].width = 22

lc = LineChart(); lc.title = "Cash balance by scenario"; lc.height = 8; lc.width = 18
lc.y_axis.title = "Cash ($)"; lc.x_axis.title = "Month"; lc.y_axis.numFmt = '$#,##0'
for si in range(4):
    col = 2 + si * 7 + 4
    ref = Reference(ws, min_col=col, min_row=first, max_row=last)
    lc.add_data(ref, titles_from_data=False)
    lc.series[-1].tx = None
from openpyxl.chart.series import SeriesLabel
for si, s in enumerate(lc.series):
    s.tx = SeriesLabel(v=scenarios[si][0])
lc.set_categories(Reference(ws, min_col=1, min_row=first, max_row=last))
lc.legend.position = "b"
ws.add_chart(lc, "J3")

# ------------------------------------------------------------- Monte Carlo
def project(growth, churn, price=inputs["price"], fixed=inputs["fixed"]):
    """growth/churn: arrays shaped (runs, months) or scalars. Returns cash and profit paths."""
    growth = np.broadcast_to(growth, (np.shape(growth)[0] if np.ndim(growth) else 1, MONTHS))
    churn = np.broadcast_to(churn, growth.shape)
    cust = inputs["customers"] * np.cumprod(1 + growth - churn, axis=1)
    rev = cust * price
    profit = rev - (rev * inputs["cogs"] + fixed)
    cash = inputs["cash"] + np.cumsum(profit, axis=1)
    return cash, profit

rng = np.random.default_rng(SEED)
# each run gets its own underlying growth and churn (how well this business really does),
# plus month-to-month noise around them
g_run = rng.normal(inputs["growth"], GROWTH_SD, (RUNS, 1))
ch_run = rng.normal(inputs["churn"], CHURN_SD, (RUNS, 1))
g = np.clip(g_run + rng.normal(0, MONTH_SD, (RUNS, MONTHS)), 0, None)
ch = np.clip(np.broadcast_to(ch_run, (RUNS, MONTHS)), 0, None)
cash, profit = project(g, ch)
pct = np.percentile(cash, [10, 25, 50, 75, 90], axis=0)
p_survive = float(np.mean((cash >= 0).all(axis=1)))
p_breakeven = float(np.mean((profit >= 0).any(axis=1)))
end_p10, end_p50, end_p90 = (float(np.percentile(cash[:, -1], q)) for q in (10, 50, 90))

sm = wb.create_sheet("Simulation")
sm["A1"] = "Monte Carlo Simulation: 1,000 Possible Futures for Example Co."; sm["A1"].font = TITLE
sm["A2"] = (f"Each run draws its own underlying growth ~ Normal({inputs['growth']:.0%}, {GROWTH_SD:.1%}) and churn ~ "
            f"Normal({inputs['churn']:.0%}, {CHURN_SD:.1%}), plus {MONTH_SD:.0%} month-to-month noise in growth; "
            f"other inputs as in the base plan. Seed {SEED}. Values are simulation output; "
            "re-run business_cashflow.py to refresh.")
sm["A2"].font = NOTE
header(sm, 4, ["Simulation result", "Value"])
res = [("Runs", RUNS, "#,##0"), ("Chance cash stays positive all 24 months", p_survive, PCT),
       ("Chance of reaching a profitable month", p_breakeven, PCT),
       ("Ending cash, bad case (10th percentile)", end_p10, USD),
       ("Ending cash, typical (median)", end_p50, USD),
       ("Ending cash, good case (90th percentile)", end_p90, USD)]
for i, (lab, v, fmt) in enumerate(res):
    sm.cell(row=5 + i, column=1, value=lab).font = BLACK
    c = sm.cell(row=5 + i, column=2, value=v); c.font = BLACK; c.number_format = fmt
sm["B6"].fill = YELLOW
header(sm, 13, ["Month", "10th pct cash ($)", "25th pct ($)", "Median ($)", "75th pct ($)", "90th pct ($)"])
for m in range(MONTHS):
    r = 14 + m
    sm.cell(row=r, column=1, value=m + 1).font = BLACK
    for k in range(5):
        c = sm.cell(row=r, column=2 + k, value=float(pct[k, m])); c.font = BLACK; c.number_format = USD
sm.column_dimensions["A"].width = 44
for col in "BCDEF":
    sm.column_dimensions[col].width = 16

mc = LineChart(); mc.title = "Simulated cash range (10th-90th percentile)"; mc.height = 8; mc.width = 18
mc.y_axis.numFmt = '$#,##0'; mc.x_axis.title = "Month"
mc.add_data(Reference(sm, min_col=2, max_col=6, min_row=13, max_row=13 + MONTHS), titles_from_data=True)
mc.set_categories(Reference(sm, min_col=1, min_row=14, max_row=13 + MONTHS))
mc.legend.position = "b"
sm.add_chart(mc, "H4")

wb.save(FILE)

# console summary for checking
for name, dp, gg, cc, dfix in scenarios:
    c_, p_ = project(np.array([[gg] * MONTHS]), np.array([[cc] * MONTHS]),
                     price=inputs["price"] * (1 + dp), fixed=inputs["fixed"] * (1 + dfix))
    neg = np.where(c_[0] < 0)[0]; be = np.where(p_[0] >= 0)[0]
    print(f"{name:20s} end cash {c_[0,-1]:>12,.0f}  m24 profit {p_[0,-1]:>9,.0f}  "
          f"runs out {neg[0]+1 if len(neg) else '-':>3}  breakeven {be[0]+1 if len(be) else '-'}")
print(f"MC: survive {p_survive:.1%}, breakeven {p_breakeven:.1%}, end P10/P50/P90 "
      f"{end_p10:,.0f} / {end_p50:,.0f} / {end_p90:,.0f}")
