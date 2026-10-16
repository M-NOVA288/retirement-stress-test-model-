"""
Practice ROI Model: how planning work turns into firm revenue.
Adds a "Practice ROI" sheet to the workbook (run after build_model.py and business_cashflow.py).

Funnel: prospects -> new clients -> retained clients -> assets under management (AUM)
Revenue: advisory fees on AUM + annuity commissions + life insurance first-year and renewal commissions
Compares a standard approach with a data-driven approach (stress test + live planner in meetings),
and values one new client with a 10-year discounted cash flow (NPV).

All assumptions are illustrative and labeled; replace them with the firm's real figures.
Run:  python practice_roi.py
"""
from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.chart import BarChart, Reference

FILE = "retirement_stress_test_model.xlsx"
F = "Arial"
BLUE = Font(name=F, color="0000FF"); BLACK = Font(name=F, color="000000"); GREEN = Font(name=F, color="008000")
BOLD = Font(name=F, bold=True); TITLE = Font(name=F, bold=True, size=14); NOTE = Font(name=F, size=9, italic=True)
HDR = Font(name=F, bold=True, color="FFFFFF"); HDR_FILL = PatternFill("solid", fgColor="1F3864")
YELLOW = PatternFill("solid", fgColor="FFFF00")
USD = '$#,##0;($#,##0);"-"'; PCT = '0.0%;(0.0%);"-"'; MULT = '0.0"x"'; NUM = '#,##0'

def header(ws, row, labels, col=1):
    for i, t in enumerate(labels):
        c = ws.cell(row=row, column=col + i, value=t)
        c.font = HDR; c.fill = HDR_FILL; c.alignment = Alignment(horizontal="center", wrap_text=True)

wb = load_workbook(FILE)
ws = wb.create_sheet("Practice ROI")
ws["A1"] = "Practice ROI Model: From Planning Conversations to Firm Revenue"; ws["A1"].font = TITLE
ws["A2"] = ("Illustrative assumptions; replace with the firm's figures. Blue = input, green = link, "
            "yellow = the four inputs the data-driven approach changes.")
ws["A2"].font = NOTE

header(ws, 4, ["Input", "Standard approach", "Data-driven approach", "Note"])
inputs = [
    # label, standard, data-driven, format, note
    ("Prospects met per month", 20, 20, NUM, "Seminars, referrals and meetings"),
    ("Conversion to client (%)", 0.15, 0.20, PCT,
     "Assumption: stress test + live planner give prospects a clear reason to act"),
    ("Annual client retention (%)", 0.92, 0.95, PCT, "Assumption: clients who understand their plan stay longer"),
    ("Average assets per new client ($)", 33000, 33000, USD,
     "Increase Consultants averages about $32.9K per client (AdvizorPro)"),
    ("New contributions per client per year ($)", 6000, 6000, USD, "Ongoing savings added to managed accounts"),
    ("Portfolio growth before fee (%)", "=Projection!$H$2+Assumptions!$B$13", "=Projection!$H$2+Assumptions!$B$13",
     PCT, "Linked: selected portfolio from the stress-test model"),
    ("Advisory fee (%/yr)", "=Assumptions!$B$13", "=Assumptions!$B$13", PCT, "Linked from Assumptions"),
    ("Clients buying an annuity (%)", 0.20, 0.25, PCT, "Assumption: stress test shows when guaranteed income helps"),
    ("Average annuity premium ($)", 50000, 50000, USD, "Assumption"),
    ("Annuity commission (% of premium)", 0.05, 0.05, PCT, "Assumption; varies by product"),
    ("Clients buying life insurance (%)", 0.25, 0.30, PCT, "Assumption"),
    ("Average annual life premium ($)", 1500, 1500, USD, "Assumption"),
    ("First-year commission (% of premium)", 0.80, 0.80, PCT, "Assumption; varies by carrier"),
    ("Renewal commission (% of premium)", 0.05, 0.05, PCT, "Assumption"),
    ("Planner cost per year ($)", 45000, 45000, USD, "Assumption: base pay plus licensing and tools"),
    ("Discount rate for client value (%)", 0.08, 0.08, PCT, "Used in the lifetime-value DCF"),
]
for i, (lab, s, d, fmt, note) in enumerate(inputs):
    r = 5 + i
    ws.cell(row=r, column=1, value=lab).font = BLACK
    for col, v in ((2, s), (3, d)):
        c = ws.cell(row=r, column=col, value=v); c.number_format = fmt
        c.font = GREEN if isinstance(v, str) else BLUE
    ws.cell(row=r, column=4, value=note).font = NOTE
for ref in ("C6", "C7", "C12", "C15"):
    ws[ref].fill = YELLOW
# rows: 5 prospects, 6 conv, 7 ret, 8 aum/client, 9 contrib, 10 growth, 11 fee, 12 ann%, 13 ann prem,
#       14 ann comm, 15 life%, 16 life prem, 17 FYC, 18 renewal, 19 cost, 20 discount

# ---------------------------------------------------------------- 3-year projection
header(ws, 22, ["Metric", "Standard Y1", "Standard Y2", "Standard Y3",
                "Data-driven Y1", "Data-driven Y2", "Data-driven Y3"])
labels = ["New clients", "Clients at year end", "Insured clients at year end", "Assets under management, year end ($)",
          "Average AUM during year ($)", "Advisory fee revenue ($)", "Annuity commissions ($)",
          "Insurance first-year commissions ($)", "Insurance renewal commissions ($)", "Total revenue ($)",
          "Planner cost ($)", "Net contribution to firm ($)", "Revenue per $1 of planner cost"]
fmts = [NUM, NUM, NUM, USD, USD, USD, USD, USD, USD, USD, USD, USD, MULT]
for i, lab in enumerate(labels):
    ws.cell(row=23 + i, column=1, value=lab).font = BOLD if lab.startswith(("Total", "Net")) else BLACK
for block, inc in ((0, "B"), (1, "C")):            # inc = input column for this scenario
    for y in range(3):
        col = 2 + block * 3 + y
        L = ws.cell(row=1, column=col).column_letter
        P = ws.cell(row=1, column=col - 1).column_letter if y > 0 else None
        prev = lambda row: f"{P}{row}" if P else "0"
        f = {
            23: f"=${inc}$5*12*${inc}$6",
            24: f"={prev(24)}*${inc}$7+{L}23",
            25: f"={prev(25)}*${inc}$7+{L}23*${inc}$15",
            26: f"={prev(26)}*${inc}$7*(1+${inc}$10)+{L}23*${inc}$8+{L}24*${inc}$9",
            27: f"=({prev(26)}+{L}26)/2",
            28: f"={L}27*${inc}$11",
            29: f"={L}23*${inc}$12*${inc}$13*${inc}$14",
            30: f"={L}23*${inc}$15*${inc}$16*${inc}$17",
            31: f"={prev(25)}*${inc}$7*${inc}$16*${inc}$18",
            32: f"=SUM({L}28:{L}31)",
            33: f"=${inc}$19",
            34: f"={L}32-{L}33",
            35: f"=IF({L}33=0,0,{L}32/{L}33)",
        }
        for r, formula in f.items():
            c = ws.cell(row=r, column=col, value=formula); c.font = BLACK; c.number_format = fmts[r - 23]

# ---------------------------------------------------------------- summary
header(ws, 37, ["3-year summary", "Standard", "Data-driven", "Difference"])
summ = [
    ("Total revenue, 3 years ($)", "=SUM(B32:D32)", "=SUM(E32:G32)", USD),
    ("Net contribution, 3 years ($)", "=SUM(B34:D34)", "=SUM(E34:G34)", USD),
    ("Clients at end of year 3", "=D24", "=G24", NUM),
    ("Assets under management, end of year 3 ($)", "=D26", "=G26", USD),
    ("Revenue per $1 of planner cost, 3 years", "=SUM(B32:D32)/SUM(B33:D33)", "=SUM(E32:G32)/SUM(E33:G33)", MULT),
]
for i, (lab, s, d, fmt) in enumerate(summ):
    r = 38 + i
    ws.cell(row=r, column=1, value=lab).font = BOLD
    for col, v in ((2, s), (3, d)):
        c = ws.cell(row=r, column=col, value=v); c.font = BLACK; c.number_format = fmt
    c = ws.cell(row=r, column=4, value=f"=C{r}-B{r}"); c.font = BOLD; c.number_format = fmt
ws["A43"] = ("Data-driven approach changes only conversion, retention, annuity and insurance take-up (yellow cells); "
             "everything else is identical, so the difference isolates their effect.")
ws["A43"].font = NOTE

# ---------------------------------------------------------------- client lifetime value DCF
header(ws, 45, ["Client lifetime value (one new client)", "Standard ($)", "Data-driven ($)"])
ws["A46"] = "Year"; ws["A46"].font = BOLD
ws["B46"] = "Expected cash flow"; ws["C46"] = "Expected cash flow"
for col in ("B46", "C46"):
    ws[col].font = BOLD
for t in range(1, 11):
    r = 46 + t
    ws.cell(row=r, column=1, value=t).font = BLACK
    for col, inc in ((2, "B"), (3, "C")):
        surv = f"${inc}$7^({t}-1)"
        aum = (f"(${inc}$8*(1+${inc}$10)^{t}+${inc}$9*((1+${inc}$10)^{t}-1)/${inc}$10)")
        fee = f"{surv}*${inc}$11*{aum}"
        if t == 1:
            extra = f"+${inc}$12*${inc}$13*${inc}$14+${inc}$15*${inc}$16*${inc}$17"
        else:
            extra = f"+{surv}*${inc}$15*${inc}$16*${inc}$18"
        c = ws.cell(row=r, column=col, value=f"={fee}{extra}"); c.font = BLACK; c.number_format = USD
ws["A57"] = "Lifetime value (NPV of 10 years)"; ws["A57"].font = BOLD
ws["B57"] = "=NPV($B$20,B47:B56)"; ws["C57"] = "=NPV($C$20,C47:C56)"
ws["A58"] = "Value of one year's new clients"; ws["A58"].font = BOLD
ws["B58"] = "=B57*B23"; ws["C58"] = "=C57*E23"
for ref in ("B57", "C57", "B58", "C58"):
    ws[ref].font = BOLD; ws[ref].number_format = USD
ws["D57"] = "=C57-B57"; ws["D58"] = "=C58-B58"
for ref in ("D57", "D58"):
    ws[ref].font = BOLD; ws[ref].number_format = USD
ws["D56"] = "Difference"; ws["D56"].font = BOLD
ws["A59"] = ("Cash flow = survival (retention ^ years) x fee on the client's growing assets, plus expected annuity and "
             "first-year insurance commissions in year 1 and insurance renewals after. Excel NPV discounts year-end flows.")
ws["A59"].font = NOTE

ws.column_dimensions["A"].width = 44
for col in "BCDEFG":
    ws.column_dimensions[col].width = 17
ws.column_dimensions["D"].width = 20

ch = BarChart(); ch.type = "col"; ch.grouping = "clustered"
ch.title = "Total revenue by year: standard vs data-driven"; ch.height = 8; ch.width = 16
ch.y_axis.numFmt = '$#,##0'
from openpyxl.chart.series import SeriesLabel
s1 = Reference(ws, min_col=2, max_col=4, min_row=32); s2 = Reference(ws, min_col=5, max_col=7, min_row=32)
ch.add_data(s1, from_rows=True, titles_from_data=False); ch.add_data(s2, from_rows=True, titles_from_data=False)
ch.series[0].tx = SeriesLabel(v="Standard"); ch.series[1].tx = SeriesLabel(v="Data-driven")
from openpyxl.chart.data_source import StrRef
ch.set_categories(Reference(ws, min_col=2, max_col=4, min_row=22))
ch.legend.position = "b"
ws.add_chart(ch, "I4")

# Dashboard link
d = wb["Dashboard"]
d["A14"] = "Practice: extra 3-year revenue (data-driven)"; d["A14"].font = BLACK
d["B14"] = "='Practice ROI'!D38"; d["B14"].font = GREEN; d["B14"].number_format = USD
d["B14"].alignment = Alignment(horizontal="right")

wb.save(FILE)
print("saved", FILE)
